from __future__ import annotations

import csv
import os
import re
import shutil
import subprocess
import sys
import threading
from pathlib import Path

from ..config import Config
from ..textutil import clean_title, is_heading_style, normalize_text
from ..utils import quick_hash
from .base import Block, ParsedDoc, ParseError


def _cell(value) -> str:
    if value is None:
        return ""
    s = str(value).replace("\r", " ").replace("\n", " ").replace("|", "\\|")
    return s[:500]


def _table_markdown(rows: list[list[str]], max_rows: int = 600) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    lines = []
    for i, row in enumerate(rows[:max_rows]):
        cells = (row + [""] * width)[:width]
        lines.append("| " + " | ".join(cells) + " |")
        if i == 0:
            lines.append("|" + "---|" * width)
    if len(rows) > max_rows:
        lines.append(f"…（共 {len(rows)} 行，已截断）")
    return "\n".join(lines)


def _docx_table(tbl) -> list[list[str]]:
    out = []
    for row in tbl.rows:
        out.append([_cell(c.text) for c in row.cells])
    return out


def _iter_docx_items(doc):
    if hasattr(doc, "iter_inner_content"):
        yield from doc.iter_inner_content()
        return
    # 兼容老版本 python-docx：先段落后表格（仅在极老版本触发）
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        yield t


def parse_docx(path: Path, cfg: Config, display_name: str | None = None,
               extra_meta: dict | None = None) -> ParsedDoc:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(str(path))
    props = doc.core_properties
    title = clean_title(props.title or "", Path(display_name or path.name).stem)
    author = (props.author or "").strip()
    blocks: list[Block] = []
    heading_stack: list[tuple[int, str]] = []
    current_path = ""
    buf: list[str] = []
    first_heading = ""

    def flush():
        nonlocal buf
        if buf:
            text = normalize_text("\n\n".join(buf))
            if text:
                blocks.append(Block(text=text, kind="text", title_path=current_path))
            buf = []

    for item in _iter_docx_items(doc):
        if isinstance(item, Paragraph):
            text = normalize_text(item.text)
            if not text:
                continue
            style_name = item.style.name if item.style is not None else ""
            is_head, level = is_heading_style(style_name)
            if is_head:
                flush()
                heading_stack = [(lv, tx) for lv, tx in heading_stack if lv < level]
                heading_stack.append((level, text))
                current_path = " > ".join(tx for _, tx in heading_stack)
                if not first_heading:
                    first_heading = text
            else:
                buf.append(text)
        elif isinstance(item, Table):
            flush()
            rows = _docx_table(item)
            md = _table_markdown(rows)
            if md:
                blocks.append(Block(text=md, kind="table", title_path=current_path))
    flush()

    if not props.title and first_heading and len(first_heading) <= 120:
        title = first_heading
    if not blocks:
        meta = {"empty": True}
        if extra_meta:
            meta.update(extra_meta)
        return ParsedDoc(title=title, author=author, page_count=0, blocks=[],
                         meta=meta, parser="docx-empty")

    meta = {
        "subject": props.subject or "",
        "keywords": props.keywords or "",
        "created": str(props.created or ""),
        "modified": str(props.modified or ""),
    }
    if extra_meta:
        meta.update(extra_meta)
    return ParsedDoc(title=title, author=author, page_count=0, blocks=blocks,
                     meta=meta, parser="docx")


def _sheet_columns(header: list, width: int) -> list[str]:
    cols = []
    for i in range(width):
        v = header[i] if i < len(header) else None
        name = str(v).strip() if v is not None and str(v).strip() else f"列{i + 1}"
        cols.append(name[:80])
    return cols


def _row_markdown(cols: list[str], rows: list[list[str]]) -> str:
    width = len(cols)
    lines = ["| " + " | ".join(_cell(c) for c in cols) + " |",
             "|" + "---|" * width]
    for row in rows:
        vals = (row + [""] * width)[:width]
        lines.append("| " + " | ".join(_cell(v) for v in vals) + " |")
    return "\n".join(lines)


def parse_xlsx(path: Path, cfg: Config, display_name: str | None = None,
               extra_meta: dict | None = None) -> ParsedDoc:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    blocks: list[Block] = []
    sheets_meta = []
    title = clean_title(wb.properties.title or "", Path(display_name or path.name).stem)
    max_text_rows = max(20, cfg.max_rows_text)
    try:
        for ws in wb.worksheets:
            try:
                max_row = int(ws.max_row or 0)
                max_col = int(ws.max_column or 0)
            except Exception:
                max_row, max_col = 0, 0
            if max_row <= 0 or max_col <= 0:
                continue
            width = min(max_col, 100)
            rows_iter = ws.iter_rows(min_row=1, max_row=min(max_row, 30),
                                     max_col=width, values_only=True)
            raw_rows = [list(r) for r in rows_iter]
            if not raw_rows:
                continue
            from ..tables import detect_header_index
            header_idx = detect_header_index(raw_rows)
            cols = _sheet_columns(raw_rows[header_idx], width)
            preamble = raw_rows[:header_idx]
            data_rows = raw_rows[header_idx + 1:]
            if len(raw_rows) >= 30 and len(data_rows) < max_text_rows:
                more = ws.iter_rows(min_row=31, max_row=min(max_row, max_text_rows + header_idx + 2),
                                    max_col=width, values_only=True)
                data_rows.extend([list(r) for r in more])
            data_rows = data_rows[:max_text_rows]
            sheets_meta.append({"name": ws.title, "rows": max_row, "cols": max_col})

            pre_text = "\n".join(
                " ".join(str(v) for v in row if v is not None and str(v).strip())
                for row in preamble if any(v is not None and str(v).strip() for v in row))
            schema = [
                f"工作表：{ws.title}",
                f"行数：{max_row}，列数：{max_col}",
                f"列名：{' | '.join(cols)}",
                *([f"表头前说明：{pre_text[:500]}"] if pre_text else []),
                f"表格样例（前{min(len(data_rows), 8)}行）：",
                _row_markdown(cols, data_rows[:8]),
            ]
            blocks.append(Block(text="\n".join(schema), kind="table",
                                title_path=f"工作表 {ws.title}", sheet=ws.title))

            # 正文按每 30 行分块，便于按内容命中具体数据
            for start in range(0, len(data_rows), 30):
                part = data_rows[start:start + 30]
                md = _row_markdown(cols, part)
                blocks.append(Block(text=f"工作表：{ws.title}\n{md}", kind="table",
                                    title_path=f"工作表 {ws.title}", sheet=ws.title))
    finally:
        try:
            wb.close()
        except Exception:
            pass

    if not blocks:
        meta = {"empty": True, "sheets": sheets_meta}
        if extra_meta:
            meta.update(extra_meta)
        return ParsedDoc(title=title, page_count=0, blocks=[], meta=meta,
                         parser="xlsx-empty")
    meta = {
        "sheets": sheets_meta,
        "subject": wb.properties.subject or "",
        "creator": wb.properties.creator or "",
    }
    if extra_meta:
        meta.update(extra_meta)
    return ParsedDoc(title=title, author=(wb.properties.creator or ""), page_count=len(sheets_meta),
                     blocks=blocks, meta=meta, parser="xlsx")


def _iter_shapes(shapes):
    for sh in shapes:
        try:
            if sh.shape_type is not None and str(sh.shape_type).startswith("GROUP"):
                yield from _iter_shapes(sh.shapes)
            else:
                yield sh
        except Exception:
            yield sh


def parse_pptx(path: Path, cfg: Config, display_name: str | None = None,
               extra_meta: dict | None = None) -> ParsedDoc:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    blocks: list[Block] = []
    title = Path(display_name or path.name).stem
    try:
        core_title = prs.core_properties.title
        if core_title:
            title = clean_title(core_title, title)
    except Exception:
        pass

    for idx, slide in enumerate(prs.slides, 1):
        slide_title = ""
        try:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                slide_title = normalize_text(slide.shapes.title.text)
        except Exception:
            pass
        parts: list[str] = []
        for sh in _iter_shapes(slide.shapes):
            try:
                if getattr(sh, "has_text_frame", False):
                    txt = normalize_text(sh.text_frame.text)
                    if txt and txt != slide_title:
                        parts.append(txt)
                elif getattr(sh, "has_table", False):
                    rows = []
                    for row in sh.table.rows:
                        rows.append([_cell(c.text) for c in row.cells])
                    md = _table_markdown(rows, max_rows=200)
                    if md:
                        parts.append(md)
                elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    alt = ""
                    try:
                        alt = sh.name or ""
                    except Exception:
                        pass
                    if alt and alt.lower() not in ("picture", "图片"):
                        parts.append(f"[图片] {alt}")
                elif getattr(sh, "has_chart", False):
                    try:
                        chart = sh.chart
                        ctitle = chart.chart_title.text_frame.text if chart.has_title else ""
                        parts.append(f"[图表] {ctitle}".strip())
                    except Exception:
                        pass
            except Exception:
                continue
        try:
            notes = slide.notes_slide.notes_text_frame.text
            if notes and notes.strip():
                parts.append("备注：" + normalize_text(notes))
        except Exception:
            pass

        text = "\n\n".join(p for p in parts if p)
        if not text:
            text = f"[空白幻灯片] {slide_title}".strip()
        blocks.append(Block(text=text, kind="slide", title_path=slide_title, slide=idx))
        if idx == 1 and slide_title:
            title = clean_title(slide_title, title)

    if not blocks:
        raise ParseError("演示文稿没有幻灯片")
    meta = {"slides": len(blocks)}
    if extra_meta:
        meta.update(extra_meta)
    return ParsedDoc(title=title, page_count=len(blocks), blocks=blocks,
                     meta=meta, parser="pptx")


_convert_lock = threading.Lock()


def _office_pids() -> set[int]:
    names = {"winword.exe", "excel.exe", "powerpnt.exe"}
    try:
        import psutil
        return {p.pid for p in psutil.process_iter(["name"])
                if (p.info.get("name") or "").lower() in names}
    except Exception:
        pass
    import subprocess
    pids: set[int] = set()
    try:
        out = subprocess.run(
            ["tasklist", "/FO", "CSV", "/NH"],
            capture_output=True, text=True, timeout=10,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        for line in (out.stdout or "").splitlines():
            parts = [p.strip('"') for p in line.split('","')]
            if len(parts) >= 2 and parts[0].lower() in names:
                try:
                    pids.add(int(parts[1]))
                except Exception:
                    pass
    except Exception:
        pass
    return pids

def _kill_pids(pids: set[int]) -> None:
    import subprocess
    for pid in pids:
        try:
            subprocess.run(["taskkill", "/PID", str(pid), "/F"],
                           capture_output=True, timeout=10,
                           creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        except Exception:
            pass


def _save_with_timeout(path: Path, out: Path, kind: str, timeout: int = 300) -> None:
    baseline = _office_pids()
    done = threading.Event()
    errors: list[Exception] = []

    def worker():
        try:
            _save_with_office(path, out, kind)
        except Exception as exc:
            errors.append(exc)
        finally:
            done.set()

    t = threading.Thread(target=worker, daemon=True)
    t.start()
    if not done.wait(timeout):
        _kill_pids(_office_pids() - baseline)
        done.wait(5)
        raise ParseError(f"Office 转换超时（{timeout} 秒），已终止卡住的 Office 进程")
    if errors:
        raise errors[0]


def _is_ole(path: Path) -> bool:
    try:
        with open(path, "rb") as f:
            return f.read(8) == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    except Exception:
        return False


def _save_with_office(path: Path, out: Path, kind: str) -> None:
    import pythoncom
    import win32com.client

    tmp = out.with_name(out.stem + f".tmp{os.getpid()}_{threading.get_ident()}{out.suffix}")
    try:
        tmp.unlink(missing_ok=True)
    except Exception:
        pass
    pythoncom.CoInitialize()
    app = None
    document = None
    try:
        if kind in ("doc", "docx"):
            app = win32com.client.DispatchEx("Word.Application")
            app.Visible = False
            app.DisplayAlerts = 0
            document = app.Documents.Open(str(path), ReadOnly=True,
                                          AddToRecentFiles=False, Visible=False)
            document.SaveAs2(str(tmp), FileFormat=16)
            document.Close(False)
            document = None
        elif kind in ("xls", "xlsx"):
            app = win32com.client.DispatchEx("Excel.Application")
            app.Visible = False
            app.DisplayAlerts = False
            document = app.Workbooks.Open(str(path), ReadOnly=True)
            document.SaveAs(str(tmp), FileFormat=51)
            document.Close(False)
            document = None
        elif kind in ("ppt", "pptx"):
            app = win32com.client.DispatchEx("PowerPoint.Application")
            document = app.Presentations.Open(str(path), ReadOnly=True, WithWindow=False)
            document.SaveAs(str(tmp), 24)
            document.Close()
            document = None
        else:
            raise ParseError(f"不支持的 Office 转换类型: {kind}")
    except Exception as exc:
        raise ParseError(f"Office 转换失败: {exc}") from exc
    finally:
        try:
            if document is not None:
                document.Close(False)
        except Exception:
            pass
        try:
            if app is not None:
                app.Quit()
        except Exception:
            pass
        try:
            del document
            del app
        except Exception:
            pass
        pythoncom.CoUninitialize()
    if not tmp.exists() or tmp.stat().st_size == 0:
        raise ParseError("Office 转换未生成有效文件")
    os.replace(tmp, out)


def _run_office_child(path: Path, out: Path, kind: str, timeout: int = 90) -> None:
    worker = Path(__file__).resolve().parent.parent / "office_convert_worker.py"
    env = dict(os.environ)
    env["PYTHONPATH"] = str(worker.parent.parent)
    env["PYTHONIOENCODING"] = "utf-8"
    cmd = [sys.executable, str(worker), str(path), str(out), kind]
    baseline = _office_pids()
    out.parent.mkdir(parents=True, exist_ok=True)
    err_path = out.with_name(out.name + f".{os.getpid()}.{threading.get_ident()}.err.txt")
    err_file = open(err_path, "w+b")
    proc = subprocess.Popen(
        cmd, cwd=str(worker.parent.parent), env=env,
        stdout=subprocess.DEVNULL, stderr=err_file,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    try:
        try:
            proc.wait(timeout=timeout)
        except subprocess.TimeoutExpired:
            subprocess.run(["taskkill", "/F", "/T", "/PID", str(proc.pid)],
                           capture_output=True,
                           creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            _kill_pids(_office_pids() - baseline)
            raise ParseError(f"Office 转换超时（{timeout} 秒）")
        if proc.returncode != 0:
            try:
                err_file.seek(0)
                msg = err_file.read().decode("utf-8", errors="ignore")[-1000:]
            except Exception:
                msg = ""
            raise ParseError(f"Office 转换失败: {msg}")
    finally:
        try:
            err_file.close()
            err_path.unlink(missing_ok=True)
        except Exception:
            pass


def convert_legacy(path: Path, cfg: Config) -> Path:
    ext = path.suffix.lower()
    target = {".doc": ".docx", ".xls": ".xlsx", ".ppt": ".pptx"}.get(ext)
    if not target:
        raise ParseError(f"不支持的历史格式: {ext}")
    out = cfg.cache_dir / "convert" / f"{quick_hash(path)}{target}"
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists() and out.stat().st_size > 0:
        return out
    _run_office_child(path, out, ext.lstrip("."))
    if not out.exists() or out.stat().st_size == 0:
        raise ParseError("Office 转换未生成有效文件")
    return out


def convert_docx_fallback(path: Path, cfg: Config) -> Path:
    out = cfg.cache_dir / "convert" / f"{quick_hash(path)}_fixed.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists() and out.stat().st_size > 0:
        return out
    _run_office_child(path, out, "docx")
    if not out.exists() or out.stat().st_size == 0:
        raise ParseError("Word 修复转换未生成有效文件")
    return out


def convert_xlsx_fallback(path: Path, cfg: Config) -> Path:
    out = cfg.cache_dir / "convert" / f"{quick_hash(path)}_fixed.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists() and out.stat().st_size > 0:
        return out
    _run_office_child(path, out, "xlsx")
    if not out.exists() or out.stat().st_size == 0:
        raise ParseError("Excel 修复转换未生成有效文件")
    return out


def _office_is_encrypted(path: Path) -> bool:
    try:
        import msoffcrypto
        with open(path, "rb") as f:
            try:
                return bool(msoffcrypto.OfficeFile(f).is_encrypted())
            except Exception:
                return False
    except Exception:
        return False


def parse_office(path: Path, cfg: Config) -> ParsedDoc:
    ext = path.suffix.lower()
    if ext in {".doc", ".xls", ".ppt", ".docx", ".xlsx", ".pptx"} and _office_is_encrypted(path):
        raise ParseError("Office 文件已加密/需要密码，按策略跳过")
    if ext in {".doc", ".xls", ".ppt"}:
        converted = convert_legacy(path, cfg)
        extra = {"source_path": str(path), "converted_path": str(converted)}
        if ext == ".doc":
            doc = parse_docx(converted, cfg, display_name=path.name, extra_meta=extra)
        elif ext == ".xls":
            doc = parse_xlsx(converted, cfg, display_name=path.name, extra_meta=extra)
        else:
            doc = parse_pptx(converted, cfg, display_name=path.name, extra_meta=extra)
        doc.parser = f"{ext[1:]}-via-office"
        return doc
    if ext == ".docx":
        return parse_docx(path, cfg)
    if ext == ".xlsx":
        return parse_xlsx(path, cfg)
    if ext == ".pptx":
        return parse_pptx(path, cfg)
    raise ParseError(f"不支持的 Office 格式: {ext}")
