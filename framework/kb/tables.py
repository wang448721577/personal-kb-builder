from __future__ import annotations

import csv
import json
import os
from datetime import date, datetime, time
from pathlib import Path

from .config import Config


def _format_value(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, time):
        return v.strftime("%H:%M:%S")
    if isinstance(v, float):
        return f"{v:.12g}"
    s = str(v)
    return s[:2000]


def detect_header_index(rows: list[list], scan: int = 15, min_cells: int = 3) -> int:
    """跳过合并标题行/制表日期行，找到真正的表头。"""
    for i, row in enumerate(rows[:scan]):
        vals = [v for v in row if v is not None and str(v).strip() != ""]
        if len(vals) >= min_cells:
            return i
    return 0


def _columns_from_header(header: list) -> list[str]:
    cols = []
    used: dict[str, int] = {}
    for i, v in enumerate(header):
        name = str(v).strip() if v is not None else ""
        name = (name or f"列{i + 1}")[:100]
        if name in used:
            used[name] += 1
            name = f"{name}_{used[name]}"
        else:
            used[name] = 1
        cols.append(name)
    return cols


def export_xlsx(path: Path, doc_id: int, cfg: Config) -> list[dict]:
    import openpyxl

    out: list[dict] = []
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        for idx, ws in enumerate(wb.worksheets):
            try:
                max_row = int(ws.max_row or 0)
                max_col = int(ws.max_column or 0)
            except Exception:
                continue
            if max_row <= 0 or max_col <= 0:
                continue
            target = cfg.table_dir / f"{doc_id}_{idx}.csv"
            tmp = target.with_suffix(".csv.tmp")
            rows_written = 0
            header: list = []
            cols: list[str] = []
            it = iter(ws.iter_rows(values_only=True))
            buffer: list[list] = []
            for row in it:
                buffer.append(list(row))
                if len(buffer) >= 15:
                    break
            header_idx = detect_header_index(buffer)
            header = buffer[header_idx] if buffer else []
            cols = _columns_from_header(header)
            with open(tmp, "w", encoding="utf-8-sig", newline="") as f:
                writer = csv.writer(f)
                if buffer:
                    writer.writerow([_format_value(v) for v in header])
                    rows_written += 1
                    for extra in buffer[header_idx + 1:]:
                        writer.writerow([_format_value(v) for v in extra])
                        rows_written += 1
                for row in it:
                    if rows_written >= cfg.max_table_rows:
                        break
                    writer.writerow([_format_value(v) for v in row])
                    rows_written += 1
            if rows_written == 0:
                tmp.unlink(missing_ok=True)
                continue
            os.replace(tmp, target)
            out.append({
                "sheet": ws.title,
                "path": str(target),
                "format": "csv",
                "rows": max(0, rows_written - 1),
                "cols": len(cols),
                "columns_json": json.dumps(cols, ensure_ascii=False),
            })
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return out


def export_csv(path: Path, doc_id: int, cfg: Config) -> list[dict]:
    target = cfg.table_dir / f"{doc_id}_0.csv"
    tmp = target.with_suffix(".csv.tmp")
    rows_written = 0
    header: list = []
    with open(path, "r", encoding="utf-8-sig", errors="ignore", newline="") as src, \
            open(tmp, "w", encoding="utf-8-sig", newline="") as dst:
        reader = csv.reader(src)
        writer = csv.writer(dst)
        for i, row in enumerate(reader):
            if i >= cfg.max_table_rows:
                break
            if i == 0:
                header = row
            writer.writerow(row)
            rows_written += 1
    if rows_written == 0:
        tmp.unlink(missing_ok=True)
        return []
    os.replace(tmp, target)
    cols = _columns_from_header(header)
    return [{
        "sheet": "CSV",
        "path": str(target),
        "format": "csv",
        "rows": max(0, rows_written - 1),
        "cols": len(cols),
        "columns_json": json.dumps(cols, ensure_ascii=False),
    }]


def export_tables(path: Path, doc_id: int, cfg: Config, ext: str | None = None) -> list[dict]:
    if not cfg.table_export:
        return []
    ext = (ext or path.suffix).lower()
    if ext == ".xlsx":
        return export_xlsx(path, doc_id, cfg)
    if ext == ".csv":
        return export_csv(path, doc_id, cfg)
    return []
